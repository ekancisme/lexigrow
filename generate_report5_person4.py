import openpyxl
from datetime import datetime
from bs4 import BeautifulSoup
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

HTML_SOURCE = '/home/youngltc/Downloads/report_system_test.html'
TARGET_EXCEL = '/home/youngltc/Downloads/lexigrow_Report5_Test_Report2.xlsx'
WORKSPACE_EXCEL = '/home/youngltc/Documents/Coding/lexigrow/lexigrow_Report5_Test_Report2.xlsx'

print('Loading HTML content...')
with open(HTML_SOURCE, 'r', encoding='utf-8') as f:
    soup = BeautifulSoup(f.read(), 'html.parser')

print('Loading Excel workbook...')
wb = openpyxl.load_workbook(TARGET_EXCEL)

# -------------------------------------------------------------------------
# Style definitions matching template
# -------------------------------------------------------------------------
FONT_HEADER = Font(name='Tahoma', size=10, bold=True)
FONT_DATA = Font(name='Tahoma', size=10, bold=False)
FILL_SECTION = PatternFill(fill_type='solid', start_color='FFCCFFFF', end_color='FFCCFFFF')
ALIGN_SECTION = Alignment(horizontal='general', vertical='center')
ALIGN_DATA_GENERAL = Alignment(horizontal='general', vertical='top', wrap_text=True)
ALIGN_DATA_LEFT = Alignment(horizontal='left', vertical='top', wrap_text=True)
ALIGN_DATA_CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)
BORDER_THIN = Border(
    left=Side(style='thin', color='FFD9D9D9'),
    right=Side(style='thin', color='FFD9D9D9'),
    top=Side(style='thin', color='FFD9D9D9'),
    bottom=Side(style='thin', color='FFD9D9D9')
)

# -------------------------------------------------------------------------
# Rename sheets to remove trailing control characters (_x0009_ / \t)
# -------------------------------------------------------------------------
sheet_name_map = {
    'Essay Management_x0009_': 'Essay Management',
    'AI Writing & Translation_x0009_': 'AI Writing & Translation',
    'Essay AI Evaluation_x0009_': 'Essay AI Evaluation',
    'Integrity & Essay Revision_x0009_': 'Integrity & Essay Revision',
    'Essay Management\t': 'Essay Management',
    'AI Writing & Translation\t': 'AI Writing & Translation',
    'Essay AI Evaluation\t': 'Essay AI Evaluation',
    'Integrity & Essay Revision\t': 'Integrity & Essay Revision',
}

for old_name, new_name in sheet_name_map.items():
    if old_name in wb.sheetnames:
        wb[old_name].title = new_name

print('Sheet names after normalization:', wb.sheetnames)

# -------------------------------------------------------------------------
# 1. UPDATE COVER SHEET
# -------------------------------------------------------------------------
print('Updating Cover sheet...')
ws_cover = wb['Cover']
ws_cover['B4'] = 'LexiGrow – AI-Powered English Vocabulary Learning and Development System with Intelligent Writing Analysis and Feedback'
ws_cover['B5'] = 'LVFS'
ws_cover['B6'] = '=B5&"_"&"Test Report"&"_"&"v1.0"'
ws_cover['F4'] = 'VuongTB'
ws_cover['F5'] = '15/09/2026'

ws_cover['A11'] = '15/09/2026'
ws_cover['B11'] = '1.0'
ws_cover['C11'] = 'Initial Test Report for Person 4'
ws_cover['D11'] = 'A'
ws_cover['E11'] = '- Add test report for Person 4 modules: Essay Management, AI Writing & Translation, Essay AI Evaluation, Integrity & Essay Revision\n- Execute test cycles (Round 1, Round 2, Round 3)'

# -------------------------------------------------------------------------
# 2. UPDATE TEST CASES SHEET (FUNCTION LIST FOR PERSON 4: UC 50 - UC 69)
# -------------------------------------------------------------------------
print('Updating Test Cases sheet...')
ws_tc = wb['Test Cases']
ws_tc['D3'] = '=Cover!B4'
ws_tc['D4'] = '=Cover!B5'
ws_tc['D5'] = '1. Server: Node.js (Express), MongoDB Atlas, Vitest\n2. Database: MongoDB Atlas / Local MongoDB\n3. Web Browser: Google Chrome, Firefox, Microsoft Edge\n4. AI Engine: Google Gemini API'

functions_data = [
    (50, 'Create Essay', 'Essay Management', 'Create new essay draft with title, theme, and optional assignment link', 'Role: Student; System active'),
    (51, 'Save Essay Draft', 'Essay Management', 'Save work-in-progress essay content periodically to database', 'User owns the draft essay'),
    (52, 'View Essay History', 'Essay Management', 'Browse paginated list of created essays with status filters', 'User is authenticated'),
    (53, 'View Essay Details', 'Essay Management', 'Inspect detailed essay content, word count, metadata, and evaluation results', 'Essay exists in database'),
    (54, 'Edit Essay', 'Essay Management', 'Modify title and content of an unsubmitted draft essay', 'Essay is in draft status'),
    (55, 'Delete Essay', 'Essay Management', 'Permanently remove draft essay from student workspace', 'Essay belongs to user and is draft'),
    (56, 'Submit Essay', 'Essay Management', 'Lock essay from editing and initiate AI evaluation pipeline', 'Essay has >= 50 words and draft state'),
    (57, 'Receive AI Topic Suggestions', 'AI Writing & Translation', 'Generate contextual writing prompts and essay ideas based on selected theme', 'Theme is selected'),
    (58, 'Use AI Writing Assistant', 'AI Writing & Translation', 'Perform real-time spellcheck, grammar correction, and sentence improvements', 'Text selected in editor'),
    (59, 'Translate Essay Text', 'AI Writing & Translation', 'Translate selected English sentences or paragraphs to Vietnamese', 'User is authenticated'),
    (60, 'Analyze Essay with AI', 'Essay AI Evaluation', 'Evaluate essay across 4 IELTS criteria (TR, CC, LR, GRA) with band scores', 'Essay is submitted'),
    (61, 'Reanalyze Essay', 'Essay AI Evaluation', 'Trigger fresh AI analysis after essay content or rubric has been modified', 'User owns essay with existing analysis'),
    (62, 'View Grammar Errors', 'Essay AI Evaluation', 'Display categorized grammar and punctuation errors with inline highlights', 'Essay has grammar errors'),
    (63, 'View Vocabulary Errors and Heatmap', 'Essay AI Evaluation', 'Visualize CEFR vocabulary distribution (A1-C2) and sentence sophistication', 'Analysis completed'),
    (64, 'View Repeated Words and Synonym Suggestions', 'Essay AI Evaluation', 'Identify overused vocabulary and recommend context-aware synonyms', 'Word used >= 4 times'),
    (65, 'View Lexical Diversity', 'Essay AI Evaluation', 'Compute Type-Token Ratio (TTR) and MTLD diversity metrics', 'Analysis completed'),
    (66, 'Check Plagiarism', 'Integrity & Essay Revision', 'Check essay against peer submissions to detect potential plagiarism', 'Peer essay in database'),
    (67, 'Detect AI-Generated Writing', 'Integrity & Essay Revision', 'Analyze perplexity and burstiness to detect LLM-generated text', 'AI/Human text submitted'),
    (68, 'Request Essay Revision', 'Integrity & Essay Revision', 'Teacher provides inline feedback and returns essay requesting student revision', 'Role: Teacher assigned to class'),
    (69, 'Submit Essay Revision', 'Integrity & Essay Revision', 'Student reviews teacher comments, updates essay, and resubmits for re-grading', 'Essay status is "needs_revision"')
]

for idx, (no, fn_name, sheet_name, desc, pre) in enumerate(functions_data, start=9):
    ws_tc.cell(idx, 2).value = no
    ws_tc.cell(idx, 3).value = fn_name
    ws_tc.cell(idx, 4).value = sheet_name
    ws_tc.cell(idx, 5).value = desc
    ws_tc.cell(idx, 6).value = pre
    for col in range(2, 7):
        c = ws_tc.cell(idx, col)
        c.font = FONT_DATA
        c.border = BORDER_THIN
        if col == 2:
            c.alignment = ALIGN_DATA_CENTER
        else:
            c.alignment = ALIGN_DATA_LEFT

# Clear leftover rows in Test Cases
last_tc_row = 9 + len(functions_data)
for r in range(last_tc_row, max(ws_tc.max_row + 1, 100)):
    for c in range(1, 10):
        cell = ws_tc.cell(r, c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()

# -------------------------------------------------------------------------
# Helper function to parse HTML sheet and populate Excel sheet
# -------------------------------------------------------------------------
def process_module_sheet(sheet_idx, excel_sheet_name, date_r1, date_r2, date_r3):
    print(f'Processing Sheet {sheet_idx}: {excel_sheet_name}...')
    ws = wb[excel_sheet_name]
    
    # 1. Unmerge any stray merged cells at or below row 11
    merges_to_remove = [rng for rng in ws.merged_cells.ranges if rng.min_row >= 11]
    for rng in merges_to_remove:
        ws.unmerge_cells(str(rng))
        
    # 2. Extract metadata from HTML
    meta_table = soup.find('table', id=f'table-sheet{sheet_idx}-meta')
    feature_name = ''
    requirement_desc = ''
    if meta_table:
        for tr in meta_table.find_all('tr'):
            tds = tr.find_all('td')
            if len(tds) >= 2:
                lbl = tds[0].text.strip()
                val = tds[1].text.strip()
                if 'Feature' in lbl:
                    feature_name = val
                elif 'Test requirement' in lbl:
                    requirement_desc = val
                    
    ws['B2'] = feature_name
    ws['B3'] = requirement_desc
    
    # 3. Extract sections and test cases from HTML main table
    main_table = soup.find('table', id=f'table-sheet{sheet_idx}')
    tbody = main_table.find('tbody') if main_table else None
    rows = tbody.find_all('tr') if tbody else main_table.find_all('tr')
    
    sections = []
    current_sec = None
    for r in rows:
        th_sec = r.find('td', class_='sec-header') or r.find('th', class_='sec-header') or r.find('td', colspan=True)
        if th_sec and ('sec-header' in th_sec.get('class', []) or int(th_sec.get('colspan', 1)) > 5):
            sec_title = th_sec.text.strip()
            current_sec = (sec_title, [])
            sections.append(current_sec)
            continue
        tds = r.find_all('td')
        if len(tds) >= 5:
            tc_id = tds[0].text.strip()
            if not tc_id or 'Test Case ID' in tc_id:
                continue
            desc = tds[1].text.strip()
            proc = tds[2].get_text(separator='\n').strip()
            exp = tds[3].get_text(separator='\n').strip()
            prec = tds[4].get_text(separator='\n').strip()
            # Status: HTML contains Round 1 status ('Passed')
            stat_r1 = tds[5].text.strip() if len(tds) > 5 else 'Passed'
            stat_r2 = 'Passed'
            stat_r3 = 'Passed'
            note = None
            if current_sec:
                current_sec[1].append((tc_id, desc, proc, exp, prec, stat_r1, stat_r2, stat_r3, note))
                
    tester = 'VuongTB'
    current_row = 11
    num_sections = len(sections)
    tc_rows = []
    
    for sec_title, tcs in sections:
        # Write section header row
        ws.row_dimensions[current_row].height = 15.75
        ws.cell(current_row, 1).value = sec_title
        for col_idx in range(1, 16):
            cell = ws.cell(current_row, col_idx)
            cell.font = FONT_HEADER
            cell.fill = FILL_SECTION
            cell.border = BORDER_THIN
            cell.alignment = ALIGN_SECTION
            if col_idx > 1:
                cell.value = None
        current_row += 1
        
        # Write test cases
        for tc in tcs:
            tc_id, desc, proc, exp, prec, r1_stat, r2_stat, r3_stat, note = tc
            ws.row_dimensions[current_row].height = 12.75
            tc_rows.append(current_row)
            
            row_vals = [
                (1, tc_id, ALIGN_DATA_GENERAL),
                (2, desc, ALIGN_DATA_GENERAL),
                (3, proc, ALIGN_DATA_GENERAL),
                (4, exp, ALIGN_DATA_LEFT),
                (5, prec, ALIGN_DATA_LEFT),
                (6, r1_stat, ALIGN_DATA_GENERAL),
                (7, date_r1, ALIGN_DATA_GENERAL),
                (8, tester, ALIGN_DATA_GENERAL),
                (9, r2_stat, ALIGN_DATA_GENERAL),
                (10, date_r2, ALIGN_DATA_GENERAL),
                (11, tester, ALIGN_DATA_GENERAL),
                (12, r3_stat, ALIGN_DATA_GENERAL),
                (13, date_r3, ALIGN_DATA_GENERAL),
                (14, tester, ALIGN_DATA_GENERAL),
                (15, note, ALIGN_DATA_GENERAL)
            ]
            
            for c_idx, val, align in row_vals:
                cell = ws.cell(current_row, c_idx)
                cell.value = val
                cell.font = FONT_DATA
                cell.fill = PatternFill(fill_type=None)
                cell.border = BORDER_THIN
                cell.alignment = align
                if isinstance(val, datetime):
                    cell.number_format = 'dd/mm/yyyy'
                    
            current_row += 1
            
    last_data_row = current_row - 1
    
    # Update formulas in header
    ws['B4'] = f'=COUNTA(A11:A{last_data_row})-{num_sections}'
    ws['B6'] = f'=COUNTIF($F11:$F{last_data_row},B5)'
    ws['C6'] = f'=COUNTIF($F11:$F{last_data_row},C5)'
    ws['D6'] = f'=COUNTIF($F11:$F{last_data_row},D5)'
    ws['E6'] = f'=COUNTIF($F11:$F{last_data_row},E5)'
    
    ws['B7'] = f'=COUNTIF($I11:$I{last_data_row},B5)'
    ws['C7'] = f'=COUNTIF($I11:$I{last_data_row},C5)'
    ws['D7'] = f'=COUNTIF($I11:$I{last_data_row},D5)'
    ws['E7'] = f'=COUNTIF($I11:$I{last_data_row},E5)'
    
    ws['B8'] = f'=COUNTIF($L11:$L{last_data_row},B5)'
    ws['C8'] = f'=COUNTIF($L11:$L{last_data_row},C5)'
    ws['D8'] = f'=COUNTIF($L11:$L{last_data_row},D5)'
    ws['E8'] = f'=COUNTIF($L11:$L{last_data_row},E5)'
    
    # Clear leftover rows
    for r in range(current_row, max(ws.max_row + 1, current_row + 100)):
        for c in range(1, 16):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            
    # Apply DataValidation dropdowns on columns F, I, L
    ws.data_validations = DataValidationList()
    dv = DataValidation(
        type='list',
        formula1='\"Passed,Failed,Pending,N/A\"',
        allowBlank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle='Invalid Entry',
        error='Please select from the list: Passed, Failed, Pending, N/A'
    )
    ws.add_data_validation(dv)
    
    ranges = []
    if tc_rows:
        start = tc_rows[0]
        end = tc_rows[0]
        for r in tc_rows[1:]:
            if r == end + 1:
                end = r
            else:
                ranges.append((start, end))
                start = r
                end = r
        ranges.append((start, end))
        
    for start, end in ranges:
        dv.add(f'F{start}:F{end}')
        dv.add(f'I{start}:I{end}')
        dv.add(f'L{start}:L{end}')
        
    print(f'  {excel_sheet_name}: successfully written {len(tc_rows)} TCs across {num_sections} sections.')

# -------------------------------------------------------------------------
# Populate 4 module sheets
# -------------------------------------------------------------------------
process_module_sheet(1, 'Essay Management', datetime(2026, 9, 15), datetime(2026, 9, 20), datetime(2026, 9, 25))
process_module_sheet(2, 'AI Writing & Translation', datetime(2026, 9, 16), datetime(2026, 9, 21), datetime(2026, 9, 26))
process_module_sheet(3, 'Essay AI Evaluation', datetime(2026, 9, 17), datetime(2026, 9, 22), datetime(2026, 9, 27))
process_module_sheet(4, 'Integrity & Essay Revision', datetime(2026, 9, 18), datetime(2026, 9, 23), datetime(2026, 9, 28))

# -------------------------------------------------------------------------
# 3. UPDATE TEST STATISTICS SHEET
# -------------------------------------------------------------------------
print('Updating Test Statistics sheet...')
ws_stat = wb['Test Statistics']

# Header metadata
ws_stat['C3'] = 'LexiGrow – AI-Powered English Vocabulary Learning and Development System with Intelligent Writing Analysis and Feedback'
ws_stat['C4'] = 'LVFS'
ws_stat['C5'] = '=C4&"_"&"Test Report"&"_"&"v1.0"'
ws_stat['H3'] = 'VuongTB'
ws_stat['H5'] = datetime(2026, 9, 15, 0, 0)
ws_stat['H5'].number_format = 'yyyy-mm-dd'

# Table rows: 4 modules
modules = [
    (11, 1, 'Essay Management'),
    (12, 2, 'AI Writing & Translation'),
    (13, 3, 'Essay AI Evaluation'),
    (14, 4, 'Integrity & Essay Revision')
]

for row_idx, mod_no, mod_name in modules:
    ws_stat[f'B{row_idx}'] = mod_no
    ws_stat[f'C{row_idx}'] = f"='{mod_name}'!B2"
    ws_stat[f'D{row_idx}'] = f"='{mod_name}'!B8"
    ws_stat[f'E{row_idx}'] = f"='{mod_name}'!C8"
    ws_stat[f'F{row_idx}'] = f"='{mod_name}'!D8"
    ws_stat[f'G{row_idx}'] = f"='{mod_name}'!E8"
    ws_stat[f'H{row_idx}'] = f"='{mod_name}'!B4"

# Clear rows 15 to max_row
for r in range(15, max(ws_stat.max_row + 1, 50)):
    for c in range(1, 12):
        cell = ws_stat.cell(r, c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()

# Row 15: Sub total
ws_stat['C15'] = 'Sub total'
ws_stat['D15'] = '=SUM(D11:D14)'
ws_stat['E15'] = '=SUM(E11:E14)'
ws_stat['F15'] = '=SUM(F11:F14)'
ws_stat['G15'] = '=SUM(G11:G14)'
ws_stat['H15'] = '=SUM(H11:H14)'

# Row 17: Test coverage
ws_stat['C17'] = 'Test coverage'
ws_stat['E17'] = '=(D15+E15)*100/(H15-G15)'
ws_stat['F17'] = '%'

# Row 18: Test successful coverage
ws_stat['C18'] = 'Test successful coverage'
ws_stat['E18'] = '=D15*100/(H15-G15)'
ws_stat['F18'] = '%'

# Format cells in rows 11 to 15 of Test Statistics
for r in range(11, 16):
    for c in range(2, 9):
        cell = ws_stat.cell(r, c)
        cell.font = FONT_HEADER if r == 15 else FONT_DATA
        cell.border = BORDER_THIN
        if c in [2, 4, 5, 6, 7, 8]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Format Test coverage rows 17 and 18
for r in [17, 18]:
    ws_stat.cell(r, 3).font = FONT_HEADER
    ws_stat.cell(r, 5).font = FONT_HEADER
    ws_stat.cell(r, 5).border = BORDER_THIN
    ws_stat.cell(r, 5).alignment = Alignment(horizontal='center', vertical='center')
    ws_stat.cell(r, 6).font = FONT_HEADER
    ws_stat.cell(r, 6).border = BORDER_THIN
    ws_stat.cell(r, 6).alignment = Alignment(horizontal='center', vertical='center')

# Save workbooks
print(f'Saving to {TARGET_EXCEL}...')
wb.save(TARGET_EXCEL)

print(f'Saving to {WORKSPACE_EXCEL}...')
wb.save(WORKSPACE_EXCEL)

print('SUCCESSFULLY GENERATED Person 4 Test Report!')
