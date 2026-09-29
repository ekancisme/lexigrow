import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
import datetime

OUTPUT_DOWNLOADS = '/home/youngltc/Downloads/1737_DN_SE_48_Report5_UnitTest.xlsx'
OUTPUT_WORKSPACE = '/home/youngltc/Documents/Coding/lexigrow/LexiGrow_16UC_UnitTest_Report.xlsx'

wb = openpyxl.load_workbook(OUTPUT_WORKSPACE)

# Styles
f_tahoma_11_bold = Font(name='Tahoma', size=11, bold=True)
f_tahoma_11_regular = Font(name='Tahoma', size=11, bold=False)
f_tahoma_8_bold_white = Font(name='Tahoma', size=8, bold=True, color='FFFFFF')
fill_navy = PatternFill(start_color='000080', end_color='000080', fill_type='solid')

border_thin_bottom = Border(bottom=Side(style='thin', color='000000'))

# 1. Format Cover Sheet
ws_cover = wb['Cover']
ws_cover.column_dimensions['A'].width = 18.0
ws_cover.column_dimensions['B'].width = 38.0
ws_cover.column_dimensions['C'].width = 20.0
ws_cover.column_dimensions['D'].width = 15.0
ws_cover.column_dimensions['E'].width = 25.0
ws_cover.column_dimensions['F'].width = 38.0

# 2. Format Functions Sheet
ws_func = wb['Functions']
ws_func.column_dimensions['A'].width = 8.0
ws_func.column_dimensions['B'].width = 36.0
ws_func.column_dimensions['C'].width = 24.0
ws_func.column_dimensions['D'].width = 42.0
ws_func.column_dimensions['E'].width = 16.0
ws_func.column_dimensions['F'].width = 16.0
ws_func.column_dimensions['G'].width = 65.0
ws_func.column_dimensions['H'].width = 55.0

ws_func.row_dimensions[10].height = 26.0
for r in range(11, 27):
    ws_func.row_dimensions[r].height = 28.0
    ws_func.cell(r, 1).alignment = Alignment(horizontal='center', vertical='center')
    ws_func.cell(r, 2).alignment = Alignment(horizontal='left', vertical='center')
    ws_func.cell(r, 3).alignment = Alignment(horizontal='left', vertical='center')
    ws_func.cell(r, 4).alignment = Alignment(horizontal='left', vertical='center')
    ws_func.cell(r, 5).alignment = Alignment(horizontal='center', vertical='center')
    ws_func.cell(r, 6).alignment = Alignment(horizontal='center', vertical='center')
    ws_func.cell(r, 7).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws_func.cell(r, 8).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

# 3. Format Statistics Sheet
ws_stat = wb['Statistics']
ws_stat.column_dimensions['A'].width = 8.0
ws_stat.column_dimensions['B'].width = 28.0
ws_stat.column_dimensions['C'].width = 14.0
ws_stat.column_dimensions['D'].width = 12.0
ws_stat.column_dimensions['E'].width = 12.0
ws_stat.column_dimensions['F'].width = 10.0
ws_stat.column_dimensions['G'].width = 10.0
ws_stat.column_dimensions['H'].width = 10.0
ws_stat.column_dimensions['I'].width = 18.0

ws_stat.row_dimensions[11].height = 26.0
for r in range(12, 28):
    ws_stat.row_dimensions[r].height = 22.0
    ws_stat.cell(r, 1).alignment = Alignment(horizontal='center', vertical='center')
    ws_stat.cell(r, 2).alignment = Alignment(horizontal='left', vertical='center')
    for c in range(3, 10):
        ws_stat.cell(r, c).alignment = Alignment(horizontal='center', vertical='center')

ws_stat.row_dimensions[28].height = 24.0
ws_stat.cell(28, 2).alignment = Alignment(horizontal='left', vertical='center')
for c in range(3, 10):
    ws_stat.cell(28, c).alignment = Alignment(horizontal='center', vertical='center')

for r in range(30, 35):
    ws_stat.row_dimensions[r].height = 22.0
    ws_stat.cell(r, 2).alignment = Alignment(horizontal='left', vertical='center')
    ws_stat.cell(r, 4).alignment = Alignment(horizontal='right', vertical='center')
    ws_stat.cell(r, 5).alignment = Alignment(horizontal='left', vertical='center')

# 4. Format all 16 Function Sheets (LVFS-01 to LVFS-16)
for i in range(1, 17):
    sheet_name = f'LVFS-{i:02d}'
    ws = wb[sheet_name]

    # Column widths
    ws.column_dimensions['A'].width = 14.0
    ws.column_dimensions['B'].width = 38.0
    ws.column_dimensions['C'].width = 4.0
    ws.column_dimensions['D'].width = 78.0
    ws.column_dimensions['E'].width = 4.0

    # Ensure all test case columns (F through T) have generous width (14.0) so dates NEVER truncate to ###
    for col_idx in range(6, 21):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 14.0

    # Header block row heights
    ws.row_dimensions[2].height = 24.0
    ws.row_dimensions[3].height = 22.0
    ws.row_dimensions[4].height = 22.0
    ws.row_dimensions[5].height = 32.0 # Test requirement
    ws.row_dimensions[6].height = 22.0
    ws.row_dimensions[7].height = 24.0
    ws.row_dimensions[9].height = 26.0 # UTCID headers
    ws.row_dimensions[10].height = 24.0

    # Test requirement text wrap
    ws['C5'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Locate Result block rows
    result_type_row = None
    pf_row = None
    date_row = None
    defect_row = None

    for r in range(35, ws.max_row + 1):
        val_a = ws.cell(r, 1).value
        val_b = ws.cell(r, 2).value
        if val_a == "Result" or (val_b and "Type(N" in str(val_b)):
            result_type_row = r
        elif val_b == "Passed/Failed":
            pf_row = r
        elif val_b == "Executed Date":
            date_row = r
        elif val_b == "Defect ID":
            defect_row = r

    # Format body rows between 11 and result_type_row - 1
    end_body = (result_type_row - 1) if result_type_row else 40
    for r in range(11, end_body + 1):
        ws.row_dimensions[r].height = 24.0
        # Col B label
        if not isinstance(ws.cell(r, 2), MergedCell) and ws.cell(r, 2).value:
            ws.cell(r, 2).alignment = Alignment(horizontal='left', vertical='center')
        # Col D description: wrap text so it never clips or overlaps
        if not isinstance(ws.cell(r, 4), MergedCell):
            ws.cell(r, 4).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        # UTCID checkmarks ('O')
        for c in range(6, 21):
            if not isinstance(ws.cell(r, c), MergedCell):
                ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center')

    # Format Result block rows
    if result_type_row:
        ws.row_dimensions[result_type_row].height = 24.0
        ws.cell(result_type_row, 1).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(result_type_row, 2).alignment = Alignment(horizontal='left', vertical='center')
        for c in range(6, 21):
            if not isinstance(ws.cell(result_type_row, c), MergedCell) and ws.cell(result_type_row, c).value:
                ws.cell(result_type_row, c).alignment = Alignment(horizontal='center', vertical='center')

    if pf_row:
        ws.row_dimensions[pf_row].height = 24.0
        ws.cell(pf_row, 2).alignment = Alignment(horizontal='left', vertical='center')
        for c in range(6, 21):
            if not isinstance(ws.cell(pf_row, c), MergedCell) and ws.cell(pf_row, c).value:
                ws.cell(pf_row, c).alignment = Alignment(horizontal='center', vertical='center')

    if date_row:
        ws.row_dimensions[date_row].height = 24.0
        ws.cell(date_row, 2).alignment = Alignment(horizontal='left', vertical='center')
        # Explicitly ensure Executed Date values are populated and formatted
        # Find how many test cases exist from row 9
        for c in range(6, 21):
            utcid_val = ws.cell(9, c).value
            if utcid_val:
                cell_date = ws.cell(date_row, c)
                cell_date.value = datetime.datetime(2026, 9, 27, 0, 0)
                cell_date.number_format = 'yyyy-mm-dd'
                cell_date.font = f_tahoma_11_regular
                cell_date.alignment = Alignment(horizontal='center', vertical='center')
                cell_date.border = border_thin_bottom

    if defect_row:
        ws.row_dimensions[defect_row].height = 22.0
        ws.cell(defect_row, 2).alignment = Alignment(horizontal='left', vertical='center')
        for c in range(6, 21):
            if not isinstance(ws.cell(defect_row, c), MergedCell):
                ws.cell(defect_row, c).border = border_thin_bottom

print('Applied all proportion and formatting adjustments.')

wb.save(OUTPUT_DOWNLOADS)
wb.save(OUTPUT_WORKSPACE)

print(f'Successfully updated and saved:\n  - {OUTPUT_DOWNLOADS}\n  - {OUTPUT_WORKSPACE}')
