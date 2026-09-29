import openpyxl
from copy import copy
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SOURCE_FILE = '/home/youngltc/Downloads/report2.xlsx'
OUTPUT_DOWNLOADS = '/home/youngltc/Downloads/report2.xlsx'
OUTPUT_WORKSPACE = '/home/youngltc/Documents/Coding/lexigrow/report2_LexiGrow_CuongLT.xlsx'

wb = openpyxl.load_workbook(SOURCE_FILE)

# -------------------------------------------------------------------------
# Style definitions matching template
# -------------------------------------------------------------------------
FONT_HEADER = Font(name='Tahoma', size=10, bold=True)
FONT_DATA = Font(name='Tahoma', size=10, bold=False)
FILL_SECTION = PatternFill(fill_type='solid', start_color='FFCCFFFF', end_color='FFCCFFFF')
ALIGN_SECTION = Alignment(horizontal='general', vertical='center')
ALIGN_DATA_GENERAL = Alignment(horizontal='general', vertical='top', wrap_text=True)
ALIGN_DATA_LEFT = Alignment(horizontal='left', vertical='top', wrap_text=True)
ALIGN_DATA_CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)
BORDER_THIN = Border(
    left=Side(style='thin', color='FFD9D9D9'),
    right=Side(style='thin', color='FFD9D9D9'),
    top=Side(style='thin', color='FFD9D9D9'),
    bottom=Side(style='thin', color='FFD9D9D9')
)

# -------------------------------------------------------------------------
# 1. UPDATE COVER SHEET
# -------------------------------------------------------------------------
print('Updating Cover sheet...')
ws_cover = wb['Cover']
ws_cover['B4'] = 'LexiGrow – AI-Powered English Vocabulary Learning and Development System with Intelligent Writing Analysis and Feedback'
ws_cover['B5'] = 'LVFS'
ws_cover['B6'] = '=B5&"_"&"Test Report"&"_"&"v1.0"'
ws_cover['F4'] = 'CuongLT'
ws_cover['F5'] = '15/09/2026'

# Record of change
ws_cover['A11'] = '15/09/2026'
ws_cover['B11'] = '1.0'
ws_cover['C11'] = 'Initial Test Report for Person 1'
ws_cover['D11'] = 'A'
ws_cover['E11'] = '- Add test report for Person 1 modules: User Authentication, User Profile & Onboarding, Subscription & Payment\n- Execute test cycles (Round 1, Round 2, Round 3)'

# -------------------------------------------------------------------------
# 2. UPDATE TEST CASES SHEET (FUNCTION LIST)
# -------------------------------------------------------------------------
print('Updating Test Cases sheet...')
ws_tc = wb['Test Cases']
ws_tc['D3'] = '=Cover!B4'
ws_tc['D4'] = '=Cover!B5'
ws_tc['D5'] = '1. Server: Node.js (Express), MongoDB Atlas, Vitest\n2. Database: MongoDB Atlas / Local MongoDB\n3. Web Browser: Google Chrome, Firefox, Microsoft Edge\n4. Payment Gateway: PayOS Sandbox Gateway'

functions_data = [
    (1, 'Register Account', 'User Authentication', 'Register new student or parent account with email OTP verification', 'Email service operational'),
    (2, 'Login', 'User Authentication', 'Authenticate user credentials and return JWT access token', 'User exists and active'),
    (3, 'Login with Google', 'User Authentication', 'Google OAuth 2.0 authorization code exchange and token generation', 'Google OAuth service operational'),
    (4, 'Verify Email', 'User Authentication', 'Verify 6-digit OTP code to activate pending user account', 'Pending user record exists'),
    (5, 'Resend Verification Email', 'User Authentication', 'Regenerate and resend 6-digit OTP verification code', 'User account in pending verification'),
    (6, 'Reset Password', 'User Authentication', 'Password reset request with OTP verification and new password update', 'User account exists in system'),
    (7, 'Get Current User Information', 'User Profile & Onboarding\t', 'Retrieve authenticated user profile, roles, and linked student accounts', 'User is authenticated with valid JWT'),
    (8, 'View Personal Profile', 'User Profile & Onboarding\t', 'Query personal profile information, CEFR level, and learning status', 'User is authenticated'),
    (9, 'Edit Personal Profile', 'User Profile & Onboarding\t', 'Update name, institution, and English proficiency level', 'User is authenticated'),
    (10, 'Change Password', 'User Profile & Onboarding\t', 'Validate current password and persist encrypted new password', 'User is authenticated with existing password'),
    (11, 'Manage Notification Preferences', 'User Profile & Onboarding\t', 'Configure email, push, and weekly summary notification preferences', 'User is authenticated'),
    (12, 'Complete Student Onboarding', 'User Profile & Onboarding\t', 'Persist initial onboarding profile or record skip status for student', 'Student is authenticated'),
    (13, 'Set Learning Preferences', 'User Profile & Onboarding\t', 'Query and update interests, target level, daily goal minutes, and pace', 'Student is authenticated'),
    (14, 'View Pricing Plans', 'Subscription & Payment\t', 'Retrieve public list of active subscription plans filterable by role', 'Pricing plans catalog available in database'),
    (15, 'Purchase Subscription', 'Subscription & Payment\t', 'Create pending payment order and generate PayOS checkout URL', 'User is authenticated; Plan is active'),
    (16, 'View Subscription and Payment History', 'Subscription & Payment\t', 'Inspect current active subscription tier and payment transactions list', 'User is authenticated')
]

for idx, (no, fn_name, sheet_name, desc, pre) in enumerate(functions_data, start=9):
    ws_tc.cell(idx, 2).value = no
    ws_tc.cell(idx, 3).value = fn_name
    ws_tc.cell(idx, 4).value = sheet_name
    ws_tc.cell(idx, 5).value = desc
    ws_tc.cell(idx, 6).value = pre
    for col in range(2, 7):
        c = ws_tc.cell(idx, col)
        c.font = FONT_DATA
        c.border = BORDER_THIN
        if col == 2:
            c.alignment = ALIGN_DATA_CENTER
        else:
            c.alignment = ALIGN_DATA_LEFT

# Clear leftover rows 25 to 100 in Test Cases
for r in range(25, 101):
    for c in range(1, 10):
        cell = ws_tc.cell(r, c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()

# -------------------------------------------------------------------------
# Helper function to write test cases into test sheets
# -------------------------------------------------------------------------
def populate_test_sheet(ws, feature_name, test_req, sections_data, sheet_tab_name):
    print(f'Populating {sheet_tab_name}...')
    ws['B2'] = feature_name
    ws['B3'] = test_req
    
    current_row = 11
    num_sections = len(sections_data)
    
    date_r1 = datetime(2026, 9, 15, 0, 0)
    date_r2 = datetime(2026, 9, 20, 0, 0)
    date_r3 = datetime(2026, 9, 25, 0, 0)
    tester = 'CuongLT'
    
    for section_title, tcs in sections_data:
        # Write Section Header Row
        ws.row_dimensions[current_row].height = 15.75
        ws.cell(current_row, 1).value = section_title
        for col_idx in range(1, 16):
            cell = ws.cell(current_row, col_idx)
            cell.font = FONT_HEADER
            cell.fill = FILL_SECTION
            cell.border = BORDER_THIN
            cell.alignment = ALIGN_SECTION
            if col_idx > 1:
                cell.value = None
        current_row += 1
        
        # Write Test Cases
        for tc in tcs:
            tc_id, desc, proc, exp, prec, r1_stat, r2_stat, r3_stat, note = tc
            ws.row_dimensions[current_row].height = 12.75
            
            row_vals = [
                (1, tc_id, ALIGN_DATA_GENERAL),
                (2, desc, ALIGN_DATA_GENERAL),
                (3, proc, ALIGN_DATA_GENERAL),
                (4, exp, ALIGN_DATA_LEFT),
                (5, prec, ALIGN_DATA_LEFT),
                (6, r1_stat, ALIGN_DATA_GENERAL),
                (7, date_r1, ALIGN_DATA_GENERAL),
                (8, tester, ALIGN_DATA_GENERAL),
                (9, r2_stat, ALIGN_DATA_GENERAL),
                (10, date_r2, ALIGN_DATA_GENERAL),
                (11, tester, ALIGN_DATA_GENERAL),
                (12, r3_stat, ALIGN_DATA_GENERAL),
                (13, date_r3, ALIGN_DATA_GENERAL),
                (14, tester, ALIGN_DATA_GENERAL),
                (15, note, ALIGN_DATA_GENERAL)
            ]
            
            for c_idx, val, align in row_vals:
                cell = ws.cell(current_row, c_idx)
                cell.value = val
                cell.font = FONT_DATA
                cell.fill = PatternFill(fill_type=None)
                cell.border = BORDER_THIN
                cell.alignment = align
                if isinstance(val, datetime):
                    cell.number_format = 'yyyy-mm-dd'
                    
            current_row += 1
            
    last_data_row = current_row - 1
    
    # Update B4 formula: Number of TCs
    ws['B4'] = f'=COUNTA(A11:A{last_data_row})-{num_sections}'
    
    # Update Round statistics formulas
    ws['B6'] = f'=COUNTIF($F11:$F{last_data_row},B5)'
    ws['C6'] = f'=COUNTIF($F11:$F{last_data_row},C5)'
    ws['D6'] = f'=COUNTIF($F11:$F{last_data_row},D5)'
    ws['E6'] = f'=COUNTIF($F11:$F{last_data_row},E5)'
    
    ws['B7'] = f'=COUNTIF($I11:$I{last_data_row},B5)'
    ws['C7'] = f'=COUNTIF($I11:$I{last_data_row},C5)'
    ws['D7'] = f'=COUNTIF($I11:$I{last_data_row},D5)'
    ws['E7'] = f'=COUNTIF($I11:$I{last_data_row},E5)'
    
    ws['B8'] = f'=COUNTIF($L11:$L{last_data_row},B5)'
    ws['C8'] = f'=COUNTIF($L11:$L{last_data_row},C5)'
    ws['D8'] = f'=COUNTIF($L11:$L{last_data_row},D5)'
    ws['E8'] = f'=COUNTIF($L11:$L{last_data_row},E5)'
    
    # Clear leftover rows
    for r in range(current_row, current_row + 150):
        for c in range(1, 16):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

# -------------------------------------------------------------------------
# DATA FOR: User Authentication
# -------------------------------------------------------------------------
auth_req = "Verifies the full user authentication flow—register, credential/Google login, email OTP verification, resend verification, and secure password reset—covering input validation, account activation, token issuance, security policies, and error responses."

auth_sections = [
    ('Register Account', [
        ('AUTH-REG-001', 'Register account with valid student information',
         '1. Navigate to Register page.\n2. Enter name: "New Student", email: "newstudent@example.com", password: "Password123", role: "student", englishLevel: "A2".\n3. Click "Register" button.',
         '1. Status Code: 201 Created.\n2. Pending user record created in database.\n3. 6-digit OTP verification email sent to user.\n4. UI redirects to OTP verification modal.',
         '1. Email service is operational.\n2. Email is not registered.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-REG-002', 'Register account with existing email (duplicate)',
         '1. Navigate to Register page.\n2. Enter registered email "student@example.com" with valid details.\n3. Click "Register" button.',
         '1. Status Code: 409 Conflict / 400 Bad Request.\n2. Error message "Email is already registered" displayed.\n3. No pending user created.',
         'User with email already exists in system.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-REG-003', 'Register account with invalid email format',
         '1. Navigate to Register page.\n2. Enter invalid email "invalid-email-format".\n3. Click "Register" button.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Please provide a valid email" displayed.\n3. Form submission blocked.',
         'Client & server validation active.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-REG-004', 'Register account with short password (< 6 chars)',
         '1. Navigate to Register page.\n2. Enter password "12345" (5 chars).\n3. Click "Register" button.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Password must be at least 6 characters" displayed.\n3. Account not created.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-REG-005', 'Register account with missing required fields',
         '1. Navigate to Register page.\n2. Leave name and email fields empty.\n3. Click "Register" button.',
         '1. Status Code: 400 Bad Request.\n2. Validation errors for required fields displayed.\n3. Form submission prevented.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-REG-006', 'Register account when email service encounters failure',
         '1. Simulate email service outage / connection timeout.\n2. Fill valid registration details.\n3. Click "Register" button.',
         '1. Status Code: 500 Internal Server Error.\n2. Error message "Email delivery failed, please try again" displayed.\n3. Pending record rolled back gracefully.',
         'Email service connection failed.', 'Failed', 'Passed', 'Passed', 'Round 1: Generic 500 error shown instead of user-friendly message.'),
        ('AUTH-REG-007', 'Register account with valid Parent role',
         '1. Navigate to Register page.\n2. Fill name: "Parent One", email: "parent@example.com", password: "Password123", role: "parent".\n3. Click "Register" button.',
         '1. Status Code: 201 Created.\n2. Pending user created with parent role.\n3. OTP verification email sent.',
         'Email service is operational.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Login', [
        ('AUTH-LOG-001', 'Login with valid email and password credentials',
         '1. Navigate to Login page.\n2. Enter email: "student@example.com", password: "Password123".\n3. Click "Login" button.',
         '1. Status Code: 200 OK.\n2. Response contains success: true and JWT access token.\n3. User details returned (role, status).\n4. Redirected to dashboard.',
         'User exists and accountStatus is active.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-LOG-002', 'Login with incorrect password',
         '1. Navigate to Login page.\n2. Enter email: "student@example.com", password: "WrongPassword!".\n3. Click "Login" button.',
         '1. Status Code: 401 Unauthorized.\n2. Error message "Invalid credentials" displayed.\n3. No JWT token issued.',
         'User exists in database.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-LOG-003', 'Login with non-existent email',
         '1. Navigate to Login page.\n2. Enter unregistered email "ghost@example.com".\n3. Click "Login" button.',
         '1. Status Code: 401 Unauthorized.\n2. Error message "Invalid credentials" displayed.\n3. No token issued.',
         'User does not exist in system.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-LOG-004', 'Login with empty email or password',
         '1. Navigate to Login page.\n2. Leave password field blank.\n3. Click "Login" button.',
         '1. Status Code: 400 Bad Request.\n2. Client/server error "Please provide email and password" displayed.\n3. Request blocked.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-LOG-005', 'Login with unverified pending account',
         '1. Navigate to Login page.\n2. Enter email of user who has not completed OTP verification.\n3. Click "Login" button.',
         '1. Status Code: 403 Forbidden / 401.\n2. Message "Please verify your email before logging in" displayed.\n3. Option to resend OTP presented.',
         'User exists in PendingUser table.', 'Failed', 'Passed', 'Passed', 'Round 1: Resend OTP link was not displayed on login failure.'),
        ('AUTH-LOG-006', 'Login with suspended or deactivated account',
         '1. Navigate to Login page.\n2. Enter credentials of account with accountStatus="suspended".\n3. Click "Login" button.',
         '1. Status Code: 403 Forbidden.\n2. Message "Account is suspended. Please contact support" displayed.\n3. No access token provided.',
         'Account status is suspended in DB.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-LOG-007', 'Rate limiting protection on multiple rapid failed logins',
         '1. Submit incorrect credentials 6 times consecutively within 1 minute.\n2. Check response on 6th attempt.',
         '1. Status Code: 429 Too Many Requests.\n2. Error message "Too many login attempts. Please try again later" displayed.\n3. Brute-force blocked.',
         'Rate limiter middleware active.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Login with Google', [
        ('AUTH-GGL-001', 'Login with valid Google OAuth code (existing user)',
         '1. Click "Sign in with Google" button.\n2. Authorize via Google consent screen.\n3. System receives valid auth code: "google-auth-code-123".\n4. Submit to POST /api/auth/google.',
         '1. Status Code: 200 OK.\n2. Google ID token verified.\n3. Existing user authenticated and JWT returned.\n4. Redirected to dashboard.',
         'Google OAuth operational; User linked with Google email.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-GGL-002', 'Login with valid Google OAuth code (new user auto-registration)',
         '1. Authenticate with Google account not previously registered.\n2. System exchanges auth code with role="student".',
         '1. Status Code: 200 OK.\n2. New User created in database with isVerified=true, accountStatus="active".\n3. Default learning profile initialized.\n4. JWT token returned.',
         'New Google user account.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-GGL-003', 'Google OAuth with missing authorization code',
         '1. Send POST request to /api/auth/google without code in body.\n2. Check response.',
         '1. Status Code: 400 Bad Request.\n2. Error message "Google authorization code is required" displayed.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-GGL-004', 'Google OAuth with expired or invalid authorization code',
         '1. Send POST to /api/auth/google with code="invalid-expired-code".',
         '1. Status Code: 401 Unauthorized / 400 Bad Request.\n2. Error message "Invalid Google authorization code" displayed.',
         'Invalid Google auth token.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-GGL-005', 'Google OAuth service timeout / network error',
         '1. Simulate network failure to Google OAuth token endpoint.\n2. Click "Sign in with Google".',
         '1. Status Code: 502 Bad Gateway / 500.\n2. Error message "Unable to connect to Google service" displayed.\n3. System fails gracefully.',
         'Google servers unreachable.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-GGL-006', 'Google OAuth with custom role assignment (Teacher)',
         '1. Authenticate via Google with role="teacher".\n2. Submit authorization request.',
         '1. Status Code: 200 OK.\n2. User account initialized with Teacher role permissions.',
         'Valid Google OAuth credentials.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Verify Email', [
        ('AUTH-VER-001', 'Verify email with correct 6-digit OTP code',
         '1. Enter email "pending@example.com" and valid OTP "654321".\n2. Click "Verify Email" button.',
         '1. Status Code: 200 OK.\n2. PendingUser deleted from database.\n3. User account created in User collection with isVerified=true.\n4. Access JWT token returned.',
         'Pending registration exists with unexpired OTP.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-VER-002', 'Verify email with incorrect OTP code',
         '1. Enter email "pending@example.com" and incorrect code "000000".\n2. Click "Verify Email" button.',
         '1. Status Code: 400 Bad Request.\n2. Error message "Invalid verification code" displayed.\n3. PendingUser status remains unchanged.',
         'Pending user exists in DB.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-VER-003', 'Verify email with expired OTP code',
         '1. Enter email with OTP code that has exceeded 10-minute validity.\n2. Click "Verify Email" button.',
         '1. Status Code: 400 Bad Request.\n2. Error message "Verification code has expired" displayed.\n3. Prompt to resend code shown.',
         'OTP code expired in DB.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-VER-004', 'Verify email for non-existent pending user',
         '1. Enter email "nobody@example.com" and code "123456".\n2. Click "Verify Email" button.',
         '1. Status Code: 404 Not Found.\n2. Error message "No pending registration found for this email" displayed.',
         'No pending record in system.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-VER-005', 'Verify email with empty code or email fields',
         '1. Leave OTP code field empty.\n2. Click "Verify Email" button.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Email and code are required" displayed.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-VER-006', 'Multiple rapid failed OTP submissions protection',
         '1. Submit wrong OTP code 5 times in rapid succession.',
         '1. Status Code: 429 Too Many Requests.\n2. Temporary cooldown enforced on OTP entry.',
         'Pending registration exists.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Resend Verification Email', [
        ('AUTH-RSN-001', 'Resend verification email for valid pending account',
         '1. Enter email "pending@example.com".\n2. Click "Resend Code" link.\n3. Check database and email dispatch.',
         '1. Status Code: 200 OK.\n2. New 6-digit OTP code generated with refreshed 10-minute expiry.\n3. New verification email sent.',
         'Pending registration exists.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RSN-002', 'Resend verification for unregistered email address',
         '1. Enter unregistered email "nonexistent@example.com".\n2. Click "Resend Code" link.',
         '1. Status Code: 404 Not Found.\n2. Error message "Pending registration not found" displayed.',
         'Email not in PendingUser.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RSN-003', 'Resend verification for already verified account',
         '1. Enter email of user who is already verified (isVerified=true).\n2. Click "Resend Code" link.',
         '1. Status Code: 400 Bad Request.\n2. Message "Account is already verified. Please login" displayed.',
         'User is verified in database.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RSN-004', 'Resend verification with empty email field',
         '1. Leave email field blank.\n2. Click "Resend Code" link.',
         '1. Status Code: 400 Bad Request.\n2. Validation message "Email is required" displayed.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RSN-005', 'Resend verification when email service fails',
         '1. Simulate mail service failure.\n2. Click "Resend Code" link.',
         '1. Status Code: 500 Internal Server Error.\n2. Error message "Failed to send email. Please try again" displayed.',
         'Mail server unavailable.', 'Failed', 'Passed', 'Passed', 'Round 1: Resend countdown timer locked button on server error.'),
    ]),
    ('Reset Password', [
        ('AUTH-RST-001', 'Step 1: Request password reset OTP with registered email',
         '1. Navigate to Forgot Password page.\n2. Enter registered email "student@example.com".\n3. Click "Send Reset Code" button.',
         '1. Status Code: 200 OK.\n2. Reset token/code generated with 10-minute expiration.\n3. Password reset email dispatched to user.',
         'User account exists in system.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RST-002', 'Step 1: Request reset OTP with unregistered email',
         '1. Enter email "unknown@example.com".\n2. Click "Send Reset Code" button.',
         '1. Status Code: 404 Not Found / 200 (secure message).\n2. Error message "User with this email not found" displayed.',
         'Email does not exist in DB.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RST-003', 'Step 1: Request reset OTP with empty email field',
         '1. Leave email field blank.\n2. Click "Send Reset Code" button.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Please provide an email" displayed.',
         'None', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RST-004', 'Step 2: Reset password with valid OTP and new password',
         '1. Enter valid email, 6-digit reset code "123456", newPassword: "BrandNewPassword123".\n2. Click "Reset Password" button.',
         '1. Status Code: 200 OK.\n2. Password successfully hashed and updated in database.\n3. Reset code cleared.\n4. User can login with new password.',
         'Valid reset code exists for user.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RST-005', 'Step 2: Reset password with incorrect OTP code',
         '1. Enter incorrect code "999999" and new password.\n2. Click "Reset Password" button.',
         '1. Status Code: 400 Bad Request.\n2. Error message "Invalid or expired reset code" displayed.\n3. Password unchanged.',
         'User exists in database.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RST-006', 'Step 2: Reset password with expired OTP code',
         '1. Enter OTP code after 10-minute expiration window.\n2. Click "Reset Password" button.',
         '1. Status Code: 400 Bad Request.\n2. Error message "Reset code has expired" displayed.',
         'Reset code expired in DB.', 'Passed', 'Passed', 'Passed', None),
        ('AUTH-RST-007', 'Step 2: Reset password with short new password (< 6 chars)',
         '1. Enter valid OTP code but newPassword: "123".\n2. Click "Reset Password" button.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Password must be at least 6 characters" displayed.',
         'Valid OTP code exists.', 'Passed', 'Passed', 'Passed', None),
    ])
]

# -------------------------------------------------------------------------
# DATA FOR: User Profile & Onboarding
# -------------------------------------------------------------------------
profile_req = "Verifies user profile retrieval and management, password changes, notification preference configurations, student onboarding completion, and personalized learning preference settings."

profile_sections = [
    ('Get Current User Information', [
        ('PROF-ME-001', 'Retrieve authenticated user info for Student role',
         '1. Send GET request to /api/auth/me.\n2. Include valid student Authorization header "Bearer valid-student-token".\n3. Verify response payload.',
         '1. Status Code: 200 OK.\n2. Response contains success: true.\n3. User details returned: id, email, name, role="student", status="active".',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-ME-002', 'Retrieve authenticated user info for Parent role (linked children)',
         '1. Send GET request to /api/auth/me with Parent token.\n2. Inspect returned user attributes.',
         '1. Status Code: 200 OK.\n2. Parent user details returned.\n3. Array of linked student children profiles included.',
         'Parent user is authenticated; Linked children exist.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-ME-003', 'Access user info without Authorization header (unauthenticated)',
         '1. Send GET request to /api/auth/me without Authorization header.\n2. Check response.',
         '1. Status Code: 401 Unauthorized.\n2. Error message "Not authorized to access this route" returned.',
         'No token provided.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-ME-004', 'Access user info with expired or malformed JWT token',
         '1. Send GET request to /api/auth/me with Authorization="Bearer expired.jwt.token".',
         '1. Status Code: 401 Unauthorized.\n2. Error message "Token is invalid or expired" returned.',
         'Invalid JWT token.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-ME-005', 'Access user info when user referenced in JWT no longer exists',
         '1. Send GET request with token for user deleted from DB.',
         '1. Status Code: 404 Not Found / 401.\n2. Error message "User account not found" returned.',
         'User record deleted in database.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('View Personal Profile', [
        ('PROF-VIEW-001', 'View personal profile with full details',
         '1. Send GET request to /api/profile with valid token.\n2. Inspect returned data object.',
         '1. Status Code: 200 OK.\n2. Profile data returned: name, email, englishLevel, institution, avatar, learningProfile.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-VIEW-002', 'View personal profile for newly registered user (default state)',
         '1. Send GET request to /api/profile for freshly activated student.\n2. Inspect profile fields.',
         '1. Status Code: 200 OK.\n2. Default values returned (englishLevel="B1", empty interests array, onboardingCompleted=false).',
         'Newly created user account.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-VIEW-003', 'View personal profile without authentication token',
         '1. Send GET request to /api/profile without Authorization header.',
         '1. Status Code: 401 Unauthorized.\n2. Error message "Access token missing" returned.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-VIEW-004', 'View personal profile with invalid authentication token',
         '1. Send GET request to /api/profile with Authorization="Bearer invalid-token".',
         '1. Status Code: 401 Unauthorized.\n2. Access denied response returned.',
         'Invalid token.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-VIEW-005', 'View personal profile when database query times out',
         '1. Simulate database timeout during profile retrieval.\n2. Send GET request to /api/profile.',
         '1. Status Code: 500 Internal Server Error.\n2. Graceful server error message returned.',
         'Database connection error.', 'Failed', 'Passed', 'Passed', 'Round 1: Timeout resulted in unhandled promise rejection.'),
    ]),
    ('Edit Personal Profile', [
        ('PROF-EDIT-001', 'Edit profile with valid comprehensive details',
         '1. Send PUT request to /api/profile with Authorization token.\n2. Body: {name: "Student Updated", englishLevel: "B2", institution: "Lexi University"}.\n3. Verify response.',
         '1. Status Code: 200 OK.\n2. Updated user profile object returned.\n3. Changes persisted correctly in User document.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-EDIT-002', 'Edit profile with partial update (name only)',
         '1. Send PUT request to /api/profile with Body: {name: "Nguyen Van A"}.\n2. Inspect other profile attributes.',
         '1. Status Code: 200 OK.\n2. Name updated to "Nguyen Van A".\n3. englishLevel and institution remain unchanged.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-EDIT-003', 'Edit profile with invalid CEFR English level',
         '1. Send PUT request to /api/profile with Body: {englishLevel: "Z9"}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Invalid English proficiency level (A1-C2 allowed)" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-EDIT-004', 'Edit profile with empty request body',
         '1. Send PUT request to /api/profile with Body: {}.',
         '1. Status Code: 400 Bad Request / 200 unchanged.\n2. Appropriate validation response returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-EDIT-005', 'Edit profile with excessively long name (> 100 chars)',
         '1. Send PUT request to /api/profile with name exceeding 100 characters.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Name cannot exceed 100 characters" returned.',
         'User is authenticated.', 'Failed', 'Passed', 'Passed', 'Round 1: String length validation not enforced on server.'),
        ('PROF-EDIT-006', 'Unauthorized edit profile attempt',
         '1. Send PUT request to /api/profile without token.',
         '1. Status Code: 401 Unauthorized.\n2. Access denied response returned.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Change Password', [
        ('PROF-PWD-001', 'Change password with correct current password and valid new password',
         '1. Send PUT request to /api/profile/password with token.\n2. Body: {currentPassword: "OldPassword123", newPassword: "NewSecretPassword456"}.\n3. Verify response and login with new password.',
         '1. Status Code: 200 OK.\n2. Success message "Password changed successfully" returned.\n3. New password encrypted and saved.\n4. Old password no longer authenticates.',
         'User is authenticated with known current password.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-PWD-002', 'Change password with incorrect current password',
         '1. Send PUT request with Body: {currentPassword: "WrongOldPassword", newPassword: "NewSecretPassword456"}.',
         '1. Status Code: 400 Bad Request / 401.\n2. Error message "Current password is incorrect" returned.\n3. Password remains unchanged.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-PWD-003', 'Change password with new password shorter than 6 characters',
         '1. Send PUT request with Body: {currentPassword: "OldPassword123", newPassword: "12345"}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "New password must be at least 6 characters" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-PWD-004', 'Change password with new password identical to current password',
         '1. Send PUT request with Body: {currentPassword: "OldPassword123", newPassword: "OldPassword123"}.',
         '1. Status Code: 400 Bad Request.\n2. Warning "New password cannot be the same as current password" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-PWD-005', 'Change password with missing required fields',
         '1. Send PUT request with Body: {currentPassword: "OldPassword123"}.\n2. Leave newPassword empty.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Current password and new password are required" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-PWD-006', 'Change password without authentication token',
         '1. Send PUT request to /api/profile/password without token.',
         '1. Status Code: 401 Unauthorized.\n2. Request rejected.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-PWD-007', 'Concurrent password change requests handling',
         '1. Send two identical change password requests simultaneously in parallel threads.',
         '1. First request returns 200 OK.\n2. Second request safely returns 400 Bad Request.\n3. Password remains consistent.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Manage Notification Preferences', [
        ('PROF-NOTIF-001', 'Update all notification preferences (email, push, weekly)',
         '1. Send PUT request to /api/profile/notifications with token.\n2. Body: {email: false, push: true, weekly: true}.\n3. Verify response.',
         '1. Status Code: 200 OK.\n2. Updated preferences returned: email=false, push=true, weekly=true.\n3. Settings saved in User document.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-NOTIF-002', 'Update partial notification preference (weekly only)',
         '1. Send PUT request with Body: {weekly: false}.\n2. Inspect other notification attributes.',
         '1. Status Code: 200 OK.\n2. weekly updated to false.\n3. Existing email and push settings remain untouched.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-NOTIF-003', 'Update notification preferences with non-boolean values',
         '1. Send PUT request with Body: {email: "not-a-bool", weekly: 12345}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Preferences must be boolean values" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-NOTIF-004', 'Update notification preferences with empty payload',
         '1. Send PUT request with Body: {}.',
         '1. Status Code: 400 Bad Request / 200 OK (no-op).\n2. System handles empty payload safely without error.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-NOTIF-005', 'Update notification preferences without authentication',
         '1. Send PUT request to /api/profile/notifications without token.',
         '1. Status Code: 401 Unauthorized.\n2. Request rejected.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('PROF-NOTIF-006', 'Database failure during preference persistence',
         '1. Simulate database write failure during notification update.',
         '1. Status Code: 500 Internal Server Error.\n2. Error message returned and original settings preserved.',
         'DB write error.', 'Failed', 'Passed', 'Passed', 'Round 1: Server returned 200 despite MongoDB write error.'),
    ]),
    ('Complete Student Onboarding', [
        ('ONBD-STU-001', 'Complete student onboarding with valid learning preferences',
         '1. Send PUT request to /api/profile/learning with Student token.\n2. Body: {onboardingCompleted: true, targetLevel: "B2", interests: ["Technology", "IELTS"], dailyGoalMinutes: 15}.\n3. Verify response.',
         '1. Status Code: 200 OK.\n2. learningProfile.onboardingCompleted set to true in database.\n3. Selected preferences persisted correctly.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('ONBD-STU-002', 'Skip onboarding step (sets onboardingCompleted=true with default values)',
         '1. Send PUT request to /api/profile/learning with Body: {onboardingCompleted: true}.\n2. Inspect learningProfile defaults.',
         '1. Status Code: 200 OK.\n2. onboardingCompleted set to true.\n3. Default level "B1" and default goal assigned.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('ONBD-STU-003', 'Non-student role attempts student onboarding',
         '1. Send PUT request to /api/profile/learning with Teacher or Parent token.\n2. Body: {onboardingCompleted: true}.',
         '1. Status Code: 403 Forbidden.\n2. Error message "User role not authorized for student onboarding" returned.',
         'User holds Teacher or Parent role.', 'Passed', 'Passed', 'Passed', None),
        ('ONBD-STU-004', 'Complete onboarding with invalid targetLevel option',
         '1. Send PUT request with Body: {onboardingCompleted: true, targetLevel: "ADVANCED_PLUS"}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error for CEFR level returned.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('ONBD-STU-005', 'Complete onboarding without authentication',
         '1. Send PUT request to /api/profile/learning without token.',
         '1. Status Code: 401 Unauthorized.\n2. Access denied.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('ONBD-STU-006', 'Re-submitting onboarding after already completed',
         '1. Submit onboarding payload for student whose onboardingCompleted is already true.\n2. Verify update behavior.',
         '1. Status Code: 200 OK.\n2. Learning profile updated successfully without error or duplication.',
         'Student onboarding already completed.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Set Learning Preferences', [
        ('PREF-LRN-001', 'Query existing learning preferences (GET)',
         '1. Send GET request to /api/profile/learning with Student token.\n2. Verify response structure.',
         '1. Status Code: 200 OK.\n2. Returns targetLevel, interests, dailyGoalMinutes, pace, timezone, and onboardingCompleted.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PREF-LRN-002', 'Update valid learning preferences (PUT)',
         '1. Send PUT request to /api/profile/learning with Student token.\n2. Body: {interests: ["Business", "IELTS"], targetLevel: "C1", dailyGoalMinutes: 20, pace: "intensive"}.\n3. Verify response.',
         '1. Status Code: 200 OK.\n2. Updated learningProfile returned with targetLevel="C1", dailyGoalMinutes=20, pace="intensive".\n3. Changes saved to DB.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PREF-LRN-003', 'Query learning preferences without authentication token',
         '1. Send GET request to /api/profile/learning without token.',
         '1. Status Code: 401 Unauthorized.\n2. Access denied.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('PREF-LRN-004', 'Update preferences with invalid targetLevel value',
         '1. Send PUT request with Body: {targetLevel: "INVALID_CEFR"}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "targetLevel must be one of A1, A2, B1, B2, C1, C2" returned.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PREF-LRN-005', 'Update preferences with negative dailyGoalMinutes',
         '1. Send PUT request with Body: {dailyGoalMinutes: -15}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "dailyGoalMinutes must be positive integer" returned.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PREF-LRN-006', 'Update preferences with invalid pace option',
         '1. Send PUT request with Body: {pace: "turbo-speed"}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "pace must be relaxed, standard, or intensive" returned.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PREF-LRN-007', 'Update preferences with empty interests list',
         '1. Send PUT request with Body: {interests: []}.\n2. Verify system handling.',
         '1. Status Code: 200 OK.\n2. Interests list updated to empty array [] safely.',
         'Student user is authenticated.', 'Passed', 'Passed', 'Passed', None),
    ])
]

# -------------------------------------------------------------------------
# DATA FOR: Subscription & Payment
# -------------------------------------------------------------------------
pay_req = "Verifies the complete subscription lifecycle—including viewing active pricing plans, purchasing subscriptions via PayOS checkout integration, webhook handling, and querying active subscription status and payment transaction history."

pay_sections = [
    ('View Pricing Plans', [
        ('PAY-PLAN-001', 'Retrieve all public active subscription plans',
         '1. Send GET request to /api/payments/plans without query parameters.\n2. Inspect returned plans array.',
         '1. Status Code: 200 OK.\n2. Returns array of active plans (Free, Plus, Pro).\n3. Each plan includes slug, name, monthlyPrice, yearlyPrice, features list.',
         'Subscription plans exist in database.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-PLAN-002', 'Filter subscription plans by target role (Student)',
         '1. Send GET request to /api/payments/plans?role=student.\n2. Inspect returned plans.',
         '1. Status Code: 200 OK.\n2. Only plans targeting Student role are returned (free-student, plus-student, pro-student).\n3. Teacher plans excluded.',
         'Student plans exist in database.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-PLAN-003', 'Filter subscription plans by target role (Teacher)',
         '1. Send GET request to /api/payments/plans?role=teacher.\n2. Inspect returned plans.',
         '1. Status Code: 200 OK.\n2. Only Teacher sponsorship and classroom plans are returned.',
         'Teacher plans exist in database.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-PLAN-004', 'Retrieve pricing plans when catalog is empty in database',
         '1. Query /api/payments/plans when no subscription plans exist in DB.\n2. Check response.',
         '1. Status Code: 200 OK.\n2. Returns empty array [].\n3. UI handles empty state gracefully.',
         'No plans in DB.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-PLAN-005', 'Retrieve pricing plans when database connection fails',
         '1. Simulate database failure during plans query.\n2. Send GET request to /api/payments/plans.',
         '1. Status Code: 500 Internal Server Error.\n2. Error message returned and server remains stable.',
         'Database connection error.', 'Failed', 'Passed', 'Passed', 'Round 1: DB connection error leaked stack trace in development mode.'),
        ('PAY-PLAN-006', 'Filter plans with non-existent role query parameter',
         '1. Send GET request to /api/payments/plans?role=nonexistent.\n2. Check response.',
         '1. Status Code: 200 OK.\n2. Returns empty array [].',
         'Plans exist in database.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('Purchase Subscription', [
        ('PAY-BUY-001', 'Purchase valid monthly Plus subscription via PayOS gateway',
         '1. Send POST request to /api/payments/create-payment-link with User token.\n2. Body: {planSlug: "plus-student", billingCycle: "monthly"}.\n3. Verify response and transaction record.',
         '1. Status Code: 200 OK.\n2. PaymentTransaction created with status="PENDING", amount=99000 VND.\n3. PayOS checkoutUrl returned (e.g., https://checkout.payos.vn/mock-checkout).\n4. orderCode and QR code returned.',
         'User is authenticated; Plus plan isActive=true; PayOS operational.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-002', 'Purchase valid yearly Pro subscription plan',
         '1. Send POST request to /api/payments/create-payment-link.\n2. Body: {planSlug: "pro-student", billingCycle: "yearly"}.\n3. Verify response.',
         '1. Status Code: 200 OK.\n2. PaymentTransaction created with amount=990000 VND (yearly discounted rate).\n3. Valid checkout URL returned.',
         'Pro plan isActive=true.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-003', 'Purchase subscription with non-existent plan slug',
         '1. Send POST request with Body: {planSlug: "non-existent-plan", billingCycle: "monthly"}.',
         '1. Status Code: 404 Not Found.\n2. Error message "Subscription plan not found" returned.\n3. No transaction created.',
         'Plan slug not in database.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-004', 'Purchase subscription with inactive plan',
         '1. Send POST request for plan with isActive=false.\n2. Body: {planSlug: "deprecated-plan", billingCycle: "monthly"}.',
         '1. Status Code: 400 Bad Request.\n2. Error message "Plan is currently inactive" returned.\n3. No order created.',
         'Plan is marked inactive.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-005', 'Purchase subscription with invalid billing cycle option',
         '1. Send POST request with Body: {planSlug: "plus-student", billingCycle: "daily"}.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Billing cycle must be monthly or yearly" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-006', 'Purchase subscription with missing required fields',
         '1. Send POST request with Body: {billingCycle: "monthly"}.\n2. Leave planSlug omitted.',
         '1. Status Code: 400 Bad Request.\n2. Validation error "Plan slug is required" returned.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-007', 'Payment transaction amount at PayOS minimum boundary (1,000 VND)',
         '1. Create payment order for plan priced at 1,000 VND minimum boundary.\n2. Submit request.',
         '1. Status Code: 200 OK.\n2. Transaction accepted and PayOS link generated successfully.',
         'Boundary amount test.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-008', 'PayOS payment gateway API timeout or error handling',
         '1. Simulate PayOS external API timeout/down during payment creation.\n2. Submit purchase order.',
         '1. Status Code: 502 Bad Gateway / 500.\n2. Error message "Payment gateway temporarily unavailable" returned.\n3. Transaction marked as FAILED or removed.',
         'PayOS gateway unavailable.', 'Failed', 'Passed', 'Passed', 'Round 1: Pending transaction was left uncleaned when PayOS API timed out.'),
        ('PAY-BUY-009', 'Purchase subscription without user authentication token',
         '1. Send POST request to /api/payments/create-payment-link without token.',
         '1. Status Code: 401 Unauthorized.\n2. Access denied.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-BUY-010', 'PayOS webhook payment confirmation processing',
         '1. Simulate PayOS webhook callback with code="00" (PAID).\n2. Send POST to /api/payments/payos-webhook with valid signature.',
         '1. Status Code: 200 OK.\n2. PaymentTransaction status updated to "PAID".\n3. User Subscription created/extended.\n4. Webhook response {success: true}.',
         'Pending transaction exists; Webhook signature valid.', 'Passed', 'Passed', 'Passed', None),
    ]),
    ('View Subscription and Payment History', [
        ('PAY-HIST-001', 'Retrieve active subscription details (User on paid Plus tier)',
         '1. Send GET request to /api/payments/my-subscription with User token.\n2. Verify returned subscription data.',
         '1. Status Code: 200 OK.\n2. Returns tier: "plus", planSlug: "plus-student", features list, and valid endDate.',
         'User has active Plus subscription.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-002', 'Retrieve subscription status for user on default Free tier',
         '1. Send GET request to /api/payments/my-subscription for user with no paid plan.\n2. Check response.',
         '1. Status Code: 200 OK.\n2. Returns tier: "free", features for free tier, and hasActiveSubscription=false.',
         'User has no paid subscription.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-003', 'Retrieve user payment transaction history with multiple records',
         '1. Send GET request to /api/payments/my-transactions with User token.\n2. Verify transactions array.',
         '1. Status Code: 200 OK.\n2. Returns array of user transactions sorted by createdAt descending.\n3. Each item includes orderCode, amount, planName, status, date.',
         'User has past payment transactions.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-004', 'Retrieve payment transaction history for user with no past orders',
         '1. Send GET request to /api/payments/my-transactions for new student.\n2. Check response.',
         '1. Status Code: 200 OK.\n2. Returns empty array [].\n3. UI displays friendly "No transaction history yet" message.',
         'No transactions in DB for user.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-005', 'Filter transactions by status (status="PAID")',
         '1. Send GET request to /api/payments/my-transactions?status=PAID.\n2. Inspect returned records.',
         '1. Status Code: 200 OK.\n2. Only transactions with status="PAID" are returned.',
         'User has transactions with multiple statuses.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-006', 'Filter transactions with invalid status query parameter',
         '1. Send GET request to /api/payments/my-transactions?status=INVALID_STATUS.\n2. Check response.',
         '1. Status Code: 400 Bad Request / 200 empty array.\n2. Request handled safely without crash.',
         'User is authenticated.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-007', 'Access subscription history without authentication token',
         '1. Send GET request to /api/payments/my-transactions without token.',
         '1. Status Code: 401 Unauthorized.\n2. Access denied.',
         'Unauthenticated request.', 'Passed', 'Passed', 'Passed', None),
        ('PAY-HIST-008', 'Database error during transaction history retrieval',
         '1. Simulate database query failure during history retrieval.\n2. Send GET request.',
         '1. Status Code: 500 Internal Server Error.\n2. Graceful error message returned.',
         'DB query failure.', 'Passed', 'Passed', 'Passed', None),
    ])
]

# Populate the 3 sheets
populate_test_sheet(wb['User Authentication'], 'User Authentication', auth_req, auth_sections, 'User Authentication')
populate_test_sheet(wb['User Profile & Onboarding\t'], 'User Profile & Onboarding', profile_req, profile_sections, 'User Profile & Onboarding\t')
populate_test_sheet(wb['Subscription & Payment\t'], 'Subscription & Payment', pay_req, pay_sections, 'Subscription & Payment\t')

# -------------------------------------------------------------------------
# 3. UPDATE TEST STATISTICS SHEET
# -------------------------------------------------------------------------
print('Updating Test Statistics sheet...')
ws_stat = wb['Test Statistics']

# Header metadata
ws_stat['C3'] = 'LexiGrow – AI-Powered English Vocabulary Learning and Development System with Intelligent Writing Analysis and Feedback'
ws_stat['C4'] = 'LVFS'
ws_stat['C5'] = '=C4&"_"&"Test Report"&"_"&"v1.0"'
ws_stat['H3'] = 'CuongLT'
ws_stat['H5'] = datetime(2026, 9, 15, 0, 0)
ws_stat['H5'].number_format = 'yyyy-mm-dd'

# Table rows: exactly 3 modules
# Row 11: User Authentication
ws_stat['B11'] = 1
ws_stat['C11'] = "='User Authentication'!B2"
ws_stat['D11'] = "='User Authentication'!B8"
ws_stat['E11'] = "='User Authentication'!C8"
ws_stat['F11'] = "='User Authentication'!D8"
ws_stat['G11'] = "='User Authentication'!E8"
ws_stat['H11'] = "='User Authentication'!B4"

# Row 12: User Profile & Onboarding
ws_stat['B12'] = 2
ws_stat['C12'] = "='User Profile & Onboarding\t'!B2"
ws_stat['D12'] = "='User Profile & Onboarding\t'!B8"
ws_stat['E12'] = "='User Profile & Onboarding\t'!C8"
ws_stat['F12'] = "='User Profile & Onboarding\t'!D8"
ws_stat['G12'] = "='User Profile & Onboarding\t'!E8"
ws_stat['H12'] = "='User Profile & Onboarding\t'!B4"

# Row 13: Subscription & Payment
ws_stat['B13'] = 3
ws_stat['C13'] = "='Subscription & Payment\t'!B2"
ws_stat['D13'] = "='Subscription & Payment\t'!B8"
ws_stat['E13'] = "='Subscription & Payment\t'!C8"
ws_stat['F13'] = "='Subscription & Payment\t'!D8"
ws_stat['G13'] = "='Subscription & Payment\t'!E8"
ws_stat['H13'] = "='Subscription & Payment\t'!B4"

# Clear rows 14 to 36 in Test Statistics
for r in range(14, 37):
    for c in range(1, 12):
        cell = ws_stat.cell(r, c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()

# Row 14: Sub total
ws_stat['C14'] = 'Sub total'
ws_stat['D14'] = '=SUM(D11:D13)'
ws_stat['E14'] = '=SUM(E11:E13)'
ws_stat['F14'] = '=SUM(F11:F13)'
ws_stat['G14'] = '=SUM(G11:G13)'
ws_stat['H14'] = '=SUM(H11:H13)'

# Row 16: Test coverage
ws_stat['C16'] = 'Test coverage'
ws_stat['E16'] = '=(D14+E14)*100/(H14-G14)'
ws_stat['F16'] = '%'

# Row 17: Test successful coverage
ws_stat['C17'] = 'Test successful coverage'
ws_stat['E17'] = '=D14*100/(H14-G14)'
ws_stat['F17'] = '%'


# Format cells in rows 11 to 14 of Test Statistics
for r in range(11, 15):
    for c in range(2, 9):
        cell = ws_stat.cell(r, c)
        cell.font = FONT_HEADER if r == 14 else FONT_DATA
        cell.border = BORDER_THIN
        if c in [2, 4, 5, 6, 7, 8]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Save workbooks
wb.save(OUTPUT_DOWNLOADS)
wb.save(OUTPUT_WORKSPACE)
print('Successfully generated report2 in Downloads and Workspace!')
