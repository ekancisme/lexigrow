import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

FILES = [
    '/home/youngltc/Downloads/report2.xlsx',
    '/home/youngltc/Documents/Coding/lexigrow/report2_LexiGrow_CuongLT.xlsx'
]

for f in FILES:
    print(f'Adding Data Validation dropdowns to {f}...')
    wb = openpyxl.load_workbook(f)
    
    for sheet_name in ['User Authentication', 'User Profile & Onboarding\t', 'Subscription & Payment\t']:
        ws = wb[sheet_name]
        
        # Reset data validations
        ws.data_validations = DataValidationList()
        
        # Create DataValidation for Passed, Failed, Pending, N/A
        dv = DataValidation(
            type='list',
            formula1='\"Passed,Failed,Pending,N/A\"',
            allowBlank=True,
            showDropDown=False,
            showErrorMessage=True,
            errorTitle='Invalid Entry',
            error='Please select from the list: Passed, Failed, Pending, N/A'
        )
        ws.add_data_validation(dv)
        
        # Identify all test case rows (rows with test case ID in Col A and description in Col B)
        tc_rows = []
        for r in range(11, ws.max_row + 1):
            val_a = ws.cell(r, 1).value
            val_b = ws.cell(r, 2).value
            if val_a is not None and val_b is not None:
                tc_rows.append(r)
                
        # Group contiguous row blocks
        ranges = []
        if tc_rows:
            start = tc_rows[0]
            end = tc_rows[0]
            for r in tc_rows[1:]:
                if r == end + 1:
                    end = r
                else:
                    ranges.append((start, end))
                    start = r
                    end = r
            ranges.append((start, end))
            
        for start, end in ranges:
            # Column F: Round 1
            dv.add(f'F{start}:F{end}')
            # Column I: Round 2
            dv.add(f'I{start}:I{end}')
            # Column L: Round 3
            dv.add(f'L{start}:L{end}')
            
        print(f'  {sheet_name}: Applied dropdown to {len(tc_rows)} rows across Round 1 (F), Round 2 (I), Round 3 (L).')
        
    wb.save(f)
    print(f'Successfully updated {f}!')

print('All data validations successfully applied to all files!')
